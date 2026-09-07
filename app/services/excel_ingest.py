"""Generic Excel (.xlsx) ingestion for per-layer test data.

Sheets arrive from testers as loosely structured "section block" documents:

    Camera testing                                  <- section title row
    Test spec name | Test spec tab name | Test Count  <- header row
    DP23_Type2     | (TS) CameraControl | 1294        <- data rows; the first
                   | (TS) Exposure ...  | 43             column is filled only
    Microscope testing                                   on the first row of a
    Test spec name | ...                                 group (forward-fill)

Column sets differ between workbooks, so nothing here is specific to any one
sheet: the parser discovers the columns from the first header row, detects
section titles, forward-fills sparse leading columns, and deduplicates rows
*within that file*. Files are never combined — two workbooks feeding the same
testing type stay two datasets, each with its own columns and history.

Row identity (`row_key`) hashes the section plus the values of the
string-typed columns only — string columns are dimensions (spec/tab names),
numeric columns are measures. A re-upload with a changed count therefore
updates the existing row instead of creating a duplicate.
"""
import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000

Cell = str | int | float | None


class IngestError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class ParsedRow:
    section: str
    values: dict[str, Cell]  # keyed by sanitized column key
    row_key: str


@dataclass
class ParsedSheet:
    """One workbook, read on its own.

    A workbook is never folded together with another: one file is one dataset
    with its own columns, its own row identity and its own snapshot timeline.
    """
    columns: list[dict]  # {key, label, type: "string"|"number"}
    sections: list[str]  # in file order
    rows: list[ParsedRow]  # unique rows, in file order
    identity_keys: list[str]  # the columns row identity is built from
    total_rows: int  # data rows before dedup
    duplicates_skipped: int

    @property
    def content_hash(self) -> str:
        """What makes reloading an unchanged workbook a no-op."""
        return content_fingerprint(self.columns, self.rows)


def sanitize_key(label: str) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
    return key or "column"


def _norm(value) -> Cell:
    """Trim/collapse whitespace (incl. NBSP), coerce numeric strings."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    if not text:
        return None
    try:
        num = float(text)
    except ValueError:
        return text
    return int(num) if num.is_integer() else num


def _shape(value: Cell) -> str:
    """A coarse pattern for a value: letter runs become A, digit runs 9, and a
    real number #. "FL-5685" and "FL-5773" share "A-9"; "9-Sep-26" and
    "7-Sep-26" share "9-A-9"; "Test spec name" is "A A A"."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return "#"
    return re.sub(r"\d+", "9", re.sub(r"[^\W\d_]+", "A", str(value)))


def _looks_like_values(cells: list[Cell], below: list[list[Cell]],
                       filled_cells: list[tuple[int, Cell]]) -> bool:
    """Is this candidate header row really the first row of data?

    A header names the values under it, so it should not look like them. When
    a column below has settled on one pattern that carries digits — an
    identifier, a date, a number — and the candidate's own cell matches that
    pattern, the candidate is a value, not a name. Two such columns, and at
    least half of them, is taken as proof.

    Digits are what make the test safe: "Name" over "Alice" shares the shape
    "A", and treating that as proof would break every table of words.
    """
    if len(below) < 2:
        return False
    strong = 0
    for column, value in filled_cells:
        seen = [row[column] for row in below
                if column < len(row) and row[column] is not None]
        if len(seen) < 2:
            continue
        pattern, count = Counter(_shape(v) for v in seen).most_common(1)[0]
        if count * 2 < len(seen) or not ("9" in pattern or "#" in pattern):
            continue  # the column has no settled, digit-bearing pattern
        if _shape(value) == pattern:
            strong += 1
    return strong >= 2 and strong * 2 >= len(filled_cells)


def _row_key(section: str, values: dict[str, Cell], identity_keys: list[str]) -> str:
    payload = json.dumps([section] + [values.get(k) for k in identity_keys],
                         ensure_ascii=False, default=str)
    return hashlib.sha256(payload.encode()).hexdigest()


def content_fingerprint(columns: list[dict], rows: list[ParsedRow]) -> str:
    """Fingerprint of what was parsed, not of the bytes it came from.

    Excel rewrites a workbook's bytes on every save — saving the same sheet
    twice without touching a cell yields two different file hashes — so bytes
    cannot answer "is this the same data?". This hashes the columns and every
    row instead, which is stable across re-saves and changes the moment a
    value does.
    """
    digest = hashlib.sha256()
    for col in columns:
        digest.update(f'{col["key"]}\x1f{col["label"]}\x1f{col["type"]}\x1e'.encode())
    for row in rows:
        payload = json.dumps([row.section, sorted(row.values.items())],
                             ensure_ascii=False, default=str, sort_keys=True)
        digest.update(payload.encode())
        digest.update(b"\x1e")
    return digest.hexdigest()


def diff_rows(previous: list[dict], current: list[ParsedRow]) -> dict:
    """Added / changed / removed between the rows of two snapshots.

    `previous` is what the database holds: dicts carrying rowKey and data.
    Only meaningful when both sides were keyed the same way — the caller
    checks that and marks the result not comparable when they were not.
    """
    before = {row["rowKey"]: row.get("data", {}) for row in previous}
    after = {row.row_key: row.values for row in current}
    added = len(after.keys() - before.keys())
    removed = len(before.keys() - after.keys())
    changed = sum(1 for key in before.keys() & after.keys() if before[key] != after[key])
    return {"added": added, "changed": changed, "removed": removed, "comparable": True}


def parse_workbook(data: bytes) -> ParsedSheet:
    if len(data) > MAX_FILE_BYTES:
        raise IngestError("file_too_large",
                          f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB limit")
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sheet = wb.worksheets[0]
        raw_rows = [[_norm(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        wb.close()
    except IngestError:
        raise
    except Exception:
        raise IngestError("unreadable_file", "File could not be read as an .xlsx workbook")

    def filled(cells: list[Cell]) -> list[tuple[int, Cell]]:
        return [(i, v) for i, v in enumerate(cells) if v is not None]

    # --- locate the header row -------------------------------------------
    # The header is the first all-text row as wide as the sheet's dominant
    # row width (so a decorative title above it, even a multi-cell one, is
    # not mistaken for the header). Fall back to the first all-text
    # multi-cell row, then to the first multi-cell row of any kind.
    #
    # Within each of those passes only the *first* eligible row is considered:
    # a header names the rows beneath it, so it comes before them. If that row
    # turns out to look like data, the pass has no header rather than a header
    # further down — otherwise the last row of a headerless sheet, which has
    # nothing under it to compare against, would be promoted instead.
    width_counts = Counter(len(filled(c)) for c in raw_rows if len(filled(c)) >= 2)
    mode_width = max(width_counts, key=lambda w: (width_counts[w], w)) if width_counts else 0

    def find_header() -> int | None:
        for want_strings, want_width in ((True, mode_width), (True, None), (False, None)):
            for i, cells in enumerate(raw_rows):
                f = filled(cells)
                if len(f) < 2 or (want_width is not None and len(f) != want_width):
                    continue
                if want_strings and not all(isinstance(v, str) for _, v in f):
                    continue
                below = [c for c in raw_rows[i + 1:] if len(filled(c)) == len(f)][:30]
                if _looks_like_values(cells, below, f):
                    break  # this pass's only candidate is data — try the next
                return i
        return None

    header_idx = find_header()
    if header_idx is None and not mode_width:
        raise IngestError("no_header_row",
                          "No header row found — the sheet needs a row of column names")

    if header_idx is None:
        # A sheet of values only. Naming the columns after the first row would
        # lose that row and label every column after whatever happened to sit
        # in it, so the columns are numbered instead and every row is data.
        first = next(c for c in raw_rows if len(filled(c)) == mode_width)
        header_start = filled(first)[0][0]
        header_labels = [f"Column {i + 1}" for i in range(mode_width)]
        data_start = 0
    else:
        header_cells = filled(raw_rows[header_idx])
        header_start = header_cells[0][0]
        header_labels = [str(v) for _, v in header_cells]
        data_start = header_idx + 1
    width = len(header_labels)
    columns = [{"key": sanitize_key(lb), "label": lb, "type": "string"} for lb in header_labels]
    # dedupe sanitized keys ("Count", "count!" -> count, count_2)
    seen: dict[str, int] = {}
    for col in columns:
        n = seen.get(col["key"], 0) + 1
        seen[col["key"]] = n
        if n > 1:
            col["key"] = f"{col['key']}_{n}"

    section = "General"
    sections: list[str] = []
    # rows above the header are titles/preamble: the first cell of the last
    # one names the opening section ("Camera testing" in the real sheets).
    # A headerless sheet has no preamble — its leading single-cell rows are
    # picked up as section titles by the data loop below.
    for cells in raw_rows[:data_start - 1] if header_idx is not None else []:
        f = filled(cells)
        if f:
            section = str(f[0][1])
            if section not in sections:
                sections.append(section)

    fill: list[Cell] = []  # previous data row, for forward-fill within a section
    data_rows: list[tuple[str, list[Cell]]] = []

    for cells in raw_rows[data_start:]:
        f = filled(cells)
        if not f:
            continue
        if len(f) == 1:
            # single non-empty cell -> section title
            section = str(f[0][1])
            if section not in sections:
                sections.append(section)
            fill = []
            continue
        # align to the header's column span
        row = list(cells[header_start:header_start + width])
        row += [None] * (width - len(row))
        present = [(j, v) for j, v in enumerate(row) if v is not None]
        if present and all(isinstance(v, str) for _, v in present):
            # repeated per-section header, possibly sloppy or partial (real
            # sheets restate the header per section, sometimes with a
            # section/device name dropped into one of the cells) — skip it
            # when at least half the columns match the header exactly
            matches = sum(v == header_labels[j] for j, v in present)
            if matches >= max(2, (width + 1) // 2):
                continue
        # forward-fill leading blanks from the previous row in this section
        for i in range(width):
            if row[i] is not None:
                break
            if i < len(fill):
                row[i] = fill[i]
        fill = row
        data_rows.append((section, row))

    if len(data_rows) > MAX_ROWS:
        raise IngestError("too_many_rows",
                          f"The sheet has {len(data_rows)} data rows (limit {MAX_ROWS})")

    keys = [c["key"] for c in columns]
    for i, col in enumerate(columns):
        col_values = [row[i] for _, row in data_rows if row[i] is not None]
        if col_values and all(isinstance(v, (int, float)) for v in col_values):
            col["type"] = "number"

    # drop separator/junk rows that carry no measure at all (every numeric
    # column empty) — but only when the sheet has numeric columns to judge by
    numeric_idx = [i for i, c in enumerate(columns) if c["type"] == "number"]
    if numeric_idx:
        data_rows = [(s, row) for s, row in data_rows
                     if any(row[i] is not None for i in numeric_idx)]
    if not data_rows:
        raise IngestError("empty_sheet", "The sheet contains no data rows")
    for row_section, _ in data_rows:
        if row_section not in sections:
            sections.append(row_section)
    sections = [s for s in sections if any(rs == s for rs, _ in data_rows)]

    identity_keys = [c["key"] for c in columns if c["type"] == "string"] or keys
    rows: list[ParsedRow] = []
    seen_keys: set[str] = set()
    duplicates = 0
    for row_section, row in data_rows:
        values = dict(zip(keys, row))
        key = _row_key(row_section, values, identity_keys)
        if key in seen_keys:
            duplicates += 1
            continue
        seen_keys.add(key)
        rows.append(ParsedRow(section=row_section, values=values, row_key=key))

    return ParsedSheet(columns=columns, sections=sections, rows=rows,
                       identity_keys=identity_keys,
                       total_rows=len(data_rows), duplicates_skipped=duplicates)
